import os
import sys

def create_presentation():
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN
        from pptx.enum.shapes import MSO_SHAPE
    except ImportError:
        print("python-pptx not yet installed. Please run: pip install python-pptx")
        return

    prs = Presentation()
    # Set 16:9 widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Colors
    DARK_BG = RGBColor(9, 13, 22)
    CYAN = RGBColor(0, 242, 254)
    TEXT_WHITE = RGBColor(248, 250, 252)
    TEXT_MUTED = RGBColor(148, 163, 184)
    ACCENT_RED = RGBColor(239, 68, 68)
    ACCENT_GREEN = RGBColor(16, 185, 129)
    CARD_BG = RGBColor(18, 28, 51)

    def set_slide_background(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = DARK_BG

    def add_header(slide, title_text, category_text="SIH 2026 | PS: SIH26184"):
        # Header box
        tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.0))
        tf = tx_box.text_frame
        tf.word_wrap = True
        
        p0 = tf.paragraphs[0]
        p0.text = category_text.upper()
        p0.font.size = Pt(11)
        p0.font.bold = True
        p0.font.color.rgb = CYAN

        p1 = tf.add_paragraph()
        p1.text = title_text
        p1.font.size = Pt(22)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_WHITE

    # ---------------- SLIDE 1: Title Slide ----------------
    slide1 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide1)

    tbox = slide1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.3), Inches(4.5))
    tf1 = tbox.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "SMART INDIA HACKATHON 2026"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = CYAN

    p = tf1.add_paragraph()
    p.text = "PROJECT PRAHAAR-AI"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE

    p = tf1.add_paragraph()
    p.text = "Predictive Real-time Analytics & Hotspot Action for Anti-fraud Recovery"
    p.font.size = Pt(18)
    p.font.color.rgb = CYAN

    p = tf1.add_paragraph()
    p.text = "\nProblem Statement ID: SIH26184 | Ministry of Home Affairs (MHA) / I4C\nTheme: Blockchain & Cybersecurity | Category: Software"
    p.font.size = Pt(13)
    p.font.color.rgb = TEXT_MUTED

    # ---------------- SLIDE 2: Problem Understanding ----------------
    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide2)
    add_header(slide2, "The Real-World Crisis: Cyber Fraud & The 'Golden Hour'")

    box = slide2.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf2 = box.text_frame
    tf2.word_wrap = True

    points2 = [
        ("The 'Golden Hour' Vulnerability:", "When victims report cyber scams to 1930 / NCRP, illicit funds are rapidly dispersed across Layer 1 -> Layer 2 -> Layer 3 mule accounts."),
        ("Fast Physical Cash Extraction:", "Within 30 to 90 minutes, organized syndicates physically withdraw cash from ATMs and CSP kiosks in remote cybercrime hubs (Mewat, Jamtara, Alwar, Deoghar)."),
        ("Flaw in Existing Paradigm:", "Current Law Enforcement systems are 100% reactive. By the time debit-freeze notices reach banks, accounts are already emptied."),
        ("Missing Intelligence Layer:", "No existing system predicts WHERE and WHEN cash will be withdrawn in advance.")
    ]

    for title, desc in points2:
        p = tf2.add_paragraph()
        p.text = f"•  {title} "
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = CYAN
        
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(12)

    # ---------------- SLIDE 3: Proposed Solution ----------------
    slide3 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide3)
    add_header(slide3, "Proposed Solution: Proactive Predictive Interception")

    box = slide3.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf3 = box.text_frame
    tf3.word_wrap = True

    points3 = [
        ("MuleGraph™ Multi-Hop Traversal:", "Constructs real-time transaction graphs to detect smurfing rings, synthetic KYC accounts, and terminal cashout nodes in <50 ms."),
        ("Spatio-Temporal Cashout Forecaster:", "Combines spatial KDE, transaction velocity, and ATM risk indices to predict target ATM/CSP coordinates 15–45 minutes in advance."),
        ("Tactical Cyber Command Center:", "Dark-mode GIS live map with pulsing danger zones, real-time 1930 stream, and 1-click incident triage."),
        ("Field Officer Tactical Mobile PWA:", "Push dispatch to on-duty PCR vans and beat patrol cops with GPS navigation and suspect profiles."),
        ("Automated Legal Compliance:", "Instant 1-click Section 106 BNSS 2023 / Section 102 CrPC debit freeze notice generation for Bank Nodal Officers.")
    ]

    for title, desc in points3:
        p = tf3.add_paragraph()
        p.text = f"✔  {title} "
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = ACCENT_GREEN
        
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(10)

    # ---------------- SLIDE 4: Technical Architecture ----------------
    slide4 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide4)
    add_header(slide4, "Technical Architecture & Data Pipeline")

    box = slide4.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf4 = box.text_frame
    tf4.word_wrap = True

    arch_text = """
    1. INGESTION LAYER:
       • Real-time WebSocket & REST connector streaming 1930 / NCRP complaint payloads.
       • Automatic data sanitization & tokenized hashing of PII data.

    2. INTELLIGENCE CORE:
       • MuleGraph Engine: Directed multi-graph analytics using NetworkX & PageRank algorithms.
       • Spatial Forecaster: Spatio-temporal model computing ATM vulnerability & cash capacity.
       • Risk Fusion Core: Generates composite confidence score & Estimated Time to Withdrawal (ETW).

    3. ACTION & INTERCEPTION LAYER:
       • Tactical Command Center Dashboard (HTML5, Leaflet GIS, Vis.js).
       • Field Beat Officer PWA for real-time mobile push dispatch.
       • Legal Freeze Engine: Digitally verifiable Section 106 BNSS / Section 102 CrPC orders.
    """
    p = tf4.paragraphs[0]
    p.text = arch_text
    p.font.size = Pt(13)
    p.font.color.rgb = TEXT_WHITE

    # ---------------- SLIDE 5: Differentiators Table ----------------
    slide5 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide5)
    add_header(slide5, "Why PRAHAAR-AI Stands Out (Competitive Advantage)")

    box = slide5.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf5 = box.text_frame
    tf5.word_wrap = True

    comp_text = """
    • RESPONSE MODEL:
      Standard Systems: Reactive investigation after cash has already left the bank.
      PRAHAAR-AI: Proactive 15–45 minute pre-withdrawal forecast during the Golden Hour.

    • MULE NETWORK DETECTION:
      Standard Systems: Single-hop account searches.
      PRAHAAR-AI: Multi-tier graph traversal (Layer 1 -> 2 -> 3) and automated ring detection.

    • GEOSPATIAL INTELLIGENCE:
      Standard Systems: Broad city/state logs.
      PRAHAAR-AI: Specific ATM Terminal ID, geocodes, and 800m confidence cordon radius.

    • LAW ENFORCEMENT DISPATCH:
      Standard Systems: Manual phone calls & delayed emails.
      PRAHAAR-AI: 1-Click direct dispatch to nearest PCR patrol vehicle & beat officer PWA.
    """
    p = tf5.paragraphs[0]
    p.text = comp_text
    p.font.size = Pt(13)
    p.font.color.rgb = TEXT_WHITE

    # Save presentation
    output_path = os.path.join(os.path.dirname(__file__), "PRAHAAR_AI_SIH26184_PitchDeck.pptx")
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    create_presentation()
